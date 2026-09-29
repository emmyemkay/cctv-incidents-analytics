from __future__ import annotations

import csv
from io import BytesIO

from django.conf import settings
from django.contrib import messages
from django.core.cache import cache
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count
from django.db.models.functions import ExtractYear
from django.http import Http404, HttpResponse, StreamingHttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone

from .forms import DatasetBulkUploadForm
from .models import DatasetUpload, Incident
from .services.analytics import (
    area_category_payload,
    category_profile_payload,
    category_time_payload,
    command_center_payload,
    dashboard_payload,
    data_quality_payload,
    dispatch_payload,
    filtered_incidents,
    analytical_category_values,
    map_report_payload,
    predictive_intelligence_payload,
    station_profile_payload,
    status_outcome_payload,
    time_patterns_payload,
)
from .services.importer import process_upload
from .services.snapshots import cached_payload, invalidate_analytics_cache
from .tasks import process_dataset_upload

EXPORT_FIELDS = [
    ("event_id", "Event ID"),
    ("producer_id", "Producer ID"),
    ("payload_version", "Payload version"),
    ("source_type", "Source"),
    ("origin", "Origin"),
    ("reported_at", "Reported at"),
    ("analytical_category", "Category"),
    ("analytical_category_key", "Category colour key"),
    ("analytical_category_colour", "Category map colour"),
    ("category", "Source category"),
    ("sub_category", "Sub-category"),
    ("case_nature", "Case nature (reference)"),
    ("incident_location", "Incident location"),
    ("governing_branch", "Governing branch"),
    ("police_station", "Police station / area"),
    ("event_status", "Status"),
    ("reporting_type", "Reporting type"),
    ("description", "Description"),
    ("processing_result", "Processing result"),
    ("feedback_content", "Feedback"),
    ("longitude", "Longitude"),
    ("latitude", "Latitude"),
    ("camera_id", "Camera ID"),
    ("detection_label", "Detection label"),
    ("confidence", "Confidence"),
]


def filter_options():
    cache_key = "analytics:filter-options:v7-gis-map-report"
    cached = cache.get(cache_key)
    if cached is not None:
        return cached
    options = {
        "sources": list(Incident.SourceType.choices),
        "years": list(
            Incident.objects.annotate(year=ExtractYear("reported_at"))
            .values_list("year", flat=True).distinct().order_by("-year")
        ),
        "statuses": list(
            Incident.objects.exclude(event_status="").values_list("event_status", flat=True)
            .distinct().order_by("event_status")
        ),
        "stations": list(
            Incident.objects.exclude(police_station="").values_list("police_station", flat=True)
            .distinct().order_by("police_station")
        ),
        "branches": list(
            Incident.objects.exclude(governing_branch="").values_list("governing_branch", flat=True)
            .distinct().order_by("governing_branch")
        ),
        "categories": analytical_category_values(),
        "sub_categories": list(
            Incident.objects.exclude(sub_category="").values_list("sub_category", flat=True)
            .distinct().order_by("sub_category")[:500]
        ),
    }
    cache.set(cache_key, options, 300)
    return options


def analytics_context(request, payload):
    return {"payload": payload, "filters": request.GET, **filter_options()}


def home(request):
    """Unified command overview and operational dashboard."""
    queryset = filtered_incidents(request.GET)
    payload = cached_payload(
        "dashboard-default", queryset, request.GET,
        lambda qs: dashboard_payload(qs, include_map=False), persist=True,
    )
    context = analytics_context(request, payload)
    context["recent_incidents"] = queryset[:12]
    return render(request, "incidents/home.html", context)


def dashboard(request):
    """Keep old bookmarks working without maintaining a duplicate dashboard."""
    return redirect("incidents:home")



def live_monitoring(request):
    params = request.GET.copy()
    params["origin"] = Incident.Origin.STREAM
    queryset = filtered_incidents(params)
    context = analytics_context(request, dashboard_payload(queryset))
    context["recent_incidents"] = queryset[:50]
    return render(request, "incidents/live_monitoring.html", context)


def _redirect_with_query(request, route_name):
    url = reverse(route_name)
    query = request.GET.urlencode()
    return redirect(f"{url}?{query}" if query else url)


def hotspot_patterns(request):
    queryset = filtered_incidents(request.GET)
    payload = {
        "area": area_category_payload(queryset),
        "category": category_time_payload(queryset),
        "time": time_patterns_payload(queryset),
    }
    return render(request, "incidents/hotspot_patterns.html", analytics_context(request, payload))


def operations_outcomes(request):
    queryset = filtered_incidents(request.GET)
    payload = {
        "command": command_center_payload(queryset),
        "dispatch": dispatch_payload(queryset),
        "status": status_outcome_payload(queryset),
    }
    return render(request, "incidents/operations_outcomes.html", analytics_context(request, payload))


def area_category_analytics(request):
    return _redirect_with_query(request, "incidents:hotspot_patterns")


def category_time_analytics(request):
    return _redirect_with_query(request, "incidents:hotspot_patterns")


def time_patterns(request):
    return _redirect_with_query(request, "incidents:hotspot_patterns")


def command_center_analysis(request):
    return _redirect_with_query(request, "incidents:operations_outcomes")


def dispatch_analysis(request):
    return _redirect_with_query(request, "incidents:operations_outcomes")


def data_quality(request):
    queryset = filtered_incidents(request.GET)
    payload = cached_payload("data-quality-default", queryset, request.GET, data_quality_payload, persist=True)
    context = analytics_context(request, payload)
    context["uploads"] = DatasetUpload.objects.all()[:20]
    return render(request, "incidents/data_quality.html", context)



def status_outcomes(request):
    return _redirect_with_query(request, "incidents:operations_outcomes")


def map_analytics_report(request):
    queryset = filtered_incidents(request.GET)
    payload = cached_payload(
        "gis-map-report", queryset, request.GET,
        lambda qs: map_report_payload(qs, include_map=False), persist=True,
    )
    context = analytics_context(request, payload)
    map_query = request.GET.copy()
    map_query.pop("category", None)
    map_query.pop("focus_category", None)
    context["map_query_without_category"] = map_query.urlencode()
    requested_focus = str(request.GET.get("focus_category", "")).strip()
    available = {row["category"] for row in payload.get("category_layers", [])}
    if requested_focus in available:
        context["initial_focus_category"] = requested_focus
    else:
        priority = payload.get("priority_layers", [])
        context["initial_focus_category"] = priority[0]["category"] if priority else ""
    return render(request, "incidents/map_analytics_report.html", context)


def location_intelligence(request):
    """Preserve the former URL while moving users to the organised GIS report."""
    return _redirect_with_query(request, "incidents:map_report")


def predictive_intelligence(request):
    queryset = filtered_incidents(request.GET)
    return render(request, "incidents/predictive_intelligence.html", analytics_context(request, predictive_intelligence_payload(queryset)))

def station_profile(request, station):
    if not Incident.objects.filter(police_station=station).exists():
        raise Http404("Unknown station or operational area")
    queryset = filtered_incidents(request.GET)
    context = analytics_context(request, station_profile_payload(queryset, station, include_map=False))
    context["profile_name"] = station
    return render(request, "incidents/station_profile.html", context)


def category_profile(request, category):
    if category not in set(analytical_category_values()):
        raise Http404("Unknown incident category")
    queryset = filtered_incidents(request.GET)
    context = analytics_context(request, category_profile_payload(queryset, category, include_map=False))
    context["profile_name"] = category
    return render(request, "incidents/category_profile.html", context)


def incident_list(request):
    queryset = filtered_incidents(request.GET)
    paginator = Paginator(queryset, 50)
    query_params = request.GET.copy()
    query_params.pop("page", None)
    context = {
        "page_obj": paginator.get_page(request.GET.get("page")),
        "filters": request.GET,
        "query_string": query_params.urlencode(),
        **filter_options(),
    }
    return render(request, "incidents/incident_list.html", context)


def _export_value(incident, field):
    value = getattr(incident, field)
    if field == "reported_at" and value:
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M:%S")
    return "" if value is None else value


class _CsvEcho:
    """File-like adapter used by csv.writer with StreamingHttpResponse."""

    @staticmethod
    def write(value):
        return value


def export_incidents_csv(request):
    queryset = filtered_incidents(request.GET).order_by("-reported_at")
    writer = csv.writer(_CsvEcho())

    def rows():
        yield writer.writerow([label for _, label in EXPORT_FIELDS])
        for incident in queryset.iterator(chunk_size=2000):
            yield writer.writerow([_export_value(incident, field) for field, _ in EXPORT_FIELDS])

    response = StreamingHttpResponse(rows(), content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="incident-export.csv"'
    response["X-Content-Type-Options"] = "nosniff"
    return response


def export_incidents_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill

    queryset = filtered_incidents(request.GET).order_by("-reported_at")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title="Incidents")
    headers = [label for _, label in EXPORT_FIELDS]

    header_row = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B5E")
        cell.alignment = Alignment(horizontal="center")
        header_row.append(cell)
    sheet.append(header_row)

    for incident in queryset.iterator(chunk_size=2000):
        sheet.append([_export_value(incident, field) for field, _ in EXPORT_FIELDS])

    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="incident-export.xlsx"'
    response["X-Content-Type-Options"] = "nosniff"
    return response


def printable_report(request):
    queryset = filtered_incidents(request.GET)
    context = analytics_context(request, dashboard_payload(queryset, include_map=False))
    context["records"] = queryset[:100]
    context["generated_at"] = timezone.localtime()
    return render(request, "incidents/print_report.html", context)


def _enqueue_upload(upload_id: int):
    task = process_dataset_upload.delay(upload_id)
    DatasetUpload.objects.filter(pk=upload_id).update(task_id=task.id or "")


def upload_dataset(request):
    if request.method == "POST":
        form = DatasetBulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            selected_files = form.cleaned_data["files"]
            requested_type = form.cleaned_data["incident_type"]
            imported_total = 0
            skipped_total = 0
            duplicate_files = 0
            failures: list[str] = []
            for uploaded_file in selected_files:
                upload = DatasetUpload.objects.create(
                    file=uploaded_file,
                    original_name=uploaded_file.name,
                    incident_type=requested_type,
                    import_method=DatasetUpload.ImportMethod.UPLOAD,
                )
                try:
                    if settings.ASYNC_DATASET_IMPORTS:
                        transaction.on_commit(lambda upload_id=upload.id: _enqueue_upload(upload_id))
                    else:
                        result = process_upload(upload)
                        imported_total += result.imported
                        skipped_total += result.skipped
                        duplicate_files += int(result.duplicate_file)
                except Exception as exc:
                    failures.append(f"{uploaded_file.name}: {exc}")
            invalidate_analytics_cache()
            if settings.ASYNC_DATASET_IMPORTS and not failures:
                messages.success(
                    request,
                    f"Queued {len(selected_files):,} dataset file(s) for background import. "
                    "Progress is available in the upload history below.",
                )
                return redirect("incidents:upload")
            if imported_total:
                messages.success(
                    request,
                    f"Imported {imported_total:,} incidents from {len(selected_files):,} file(s); "
                    f"skipped {skipped_total:,} duplicate or invalid rows.",
                )
            elif duplicate_files == len(selected_files):
                messages.info(request, "All selected files had already been imported. No duplicate incidents were created.")
            else:
                messages.info(request, f"No new incidents were added; {skipped_total:,} rows were already present or invalid.")
            for failure in failures:
                messages.error(request, f"Import failed — {failure}")
            return redirect("incidents:upload" if failures else "incidents:home")
    else:
        form = DatasetBulkUploadForm()

    source_counts = {
        row["source_type"]: row["count"]
        for row in Incident.objects.values("source_type").annotate(count=Count("id"))
    }
    return render(
        request,
        "incidents/upload.html",
        {
            "form": form,
            "uploads": DatasetUpload.objects.all()[:30],
            "incident_total": Incident.objects.count(),
            "fire_count": source_counts.get(Incident.SourceType.FIRE, 0),
            "crime_count": source_counts.get(Incident.SourceType.CRIME, 0),
            "traffic_count": source_counts.get(Incident.SourceType.TRAFFIC, 0),
            "cctv_count": source_counts.get(Incident.SourceType.CCTV, 0),
            "has_active_imports": DatasetUpload.objects.filter(
                status__in=[DatasetUpload.Status.PENDING, DatasetUpload.Status.PROCESSING]
            ).exists(),
        },
    )
